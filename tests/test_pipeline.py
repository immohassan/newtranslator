"""End-to-end document pipelines: text translated, layout mirrored, style kept."""
import os

import fitz
import pytest
from docx import Document as DocxFile
from docx.oxml.ns import qn

from app.core.extract import extract_pdf
from app.core.pipeline import TranslationOptions, run_pipeline
from app.core.qa import QAReport
from app.core.shape_arabic import contains_arabic, is_already_shaped


def page_text(path):
    with fitz.open(path) as pdf:
        return pdf[0].get_text()


def page_blocks(path):
    with fitz.open(path) as pdf:
        page = pdf[0]
        out = []
        for block in page.get_text("dict")["blocks"]:
            if block.get("type") != 0:
                continue
            text = "".join(s["text"] for l in block["lines"] for s in l["spans"])
            size = block["lines"][0]["spans"][0]["size"]
            out.append({"text": text, "bbox": block["bbox"], "size": size})
        return out, page.rect.width


# ---------------------------------------------------------------- PDF
def test_pdf_en2ar_translates_text(sample_pdf, tmp_path):
    out = str(tmp_path / "out.pdf")
    run_pipeline(sample_pdf, out, TranslationOptions("en2ar", mirror=True, html_engine=False))
    text = page_text(out)
    assert contains_arabic(text), "output must contain Arabic"
    assert "Quarterly Report" not in text, "source text must be replaced"


def test_pdf_arabic_uses_joined_letterforms(sample_pdf, tmp_path):
    """The whole point of arabic_reshaper: PDF text must be presentation forms,
    not raw isolated letters."""
    out = str(tmp_path / "out.pdf")
    run_pipeline(sample_pdf, out, TranslationOptions("en2ar", mirror=True, html_engine=False))
    assert is_already_shaped(page_text(out))


def test_pdf_no_missing_glyphs(sample_pdf, tmp_path):
    """A null byte in the extracted text means the font lacked the glyph."""
    out = str(tmp_path / "out.pdf")
    run_pipeline(sample_pdf, out, TranslationOptions("en2ar", mirror=True, html_engine=False))
    assert "\x00" not in page_text(out)


def test_pdf_mirror_moves_image_to_other_side(sample_pdf, tmp_path):
    out = str(tmp_path / "out.pdf")
    run_pipeline(sample_pdf, out, TranslationOptions("en2ar", mirror=True, html_engine=False))

    before = extract_pdf(sample_pdf, QAReport()).pages[0]
    original_x0 = before.images[0].bbox.x0
    width = before.width

    with fitz.open(out) as pdf:
        page = pdf[0]
        rect = page.get_image_rects(page.get_images(full=True)[0][0])[0]

    assert original_x0 > width / 2, "fixture arrow starts on the right"
    assert rect.x0 < width / 2, "after mirroring it must be on the left"
    assert rect.x0 == pytest.approx(width - before.images[0].bbox.x1, abs=1)


def test_pdf_text_right_aligned_after_mirror(sample_pdf, tmp_path):
    out = str(tmp_path / "out.pdf")
    run_pipeline(sample_pdf, out, TranslationOptions("en2ar", mirror=True, html_engine=False))
    blocks, width = page_blocks(out)
    arabic = [b for b in blocks if contains_arabic(b["text"])]
    assert arabic
    for block in arabic:
        assert block["bbox"][2] > width / 2, "Arabic must sit on the right"


def test_pdf_nothing_drawn_off_page(sample_pdf, tmp_path):
    out = str(tmp_path / "out.pdf")
    run_pipeline(sample_pdf, out, TranslationOptions("en2ar", mirror=True, html_engine=False))
    blocks, width = page_blocks(out)
    for block in blocks:
        assert block["bbox"][0] >= -1
        assert block["bbox"][2] <= width + 1, f"{block['text'][:20]} runs off the page"


def test_pdf_headline_stays_larger_than_body(sample_pdf, tmp_path):
    """Relative sizing must survive translation."""
    out = str(tmp_path / "out.pdf")
    run_pipeline(sample_pdf, out, TranslationOptions("en2ar", mirror=True, html_engine=False))
    blocks, _ = page_blocks(out)
    sizes = sorted(b["size"] for b in blocks)
    assert sizes[-1] > sizes[0] * 1.5, "the headline must remain visibly bigger"


def test_pdf_no_mirror_keeps_positions(sample_pdf, tmp_path):
    out = str(tmp_path / "out.pdf")
    run_pipeline(sample_pdf, out, TranslationOptions("en2ar", mirror=False, html_engine=False))
    before = extract_pdf(sample_pdf, QAReport()).pages[0]
    with fitz.open(out) as pdf:
        page = pdf[0]
        rect = page.get_image_rects(page.get_images(full=True)[0][0])[0]
    assert rect.x0 == pytest.approx(before.images[0].bbox.x0, abs=1)


def test_pdf_output_is_still_a_pdf(sample_pdf, tmp_path):
    out = str(tmp_path / "out.pdf")
    run_pipeline(sample_pdf, out, TranslationOptions("en2ar", html_engine=False))
    with open(out, "rb") as fh:
        assert fh.read(5) == b"%PDF-", "the file must stay a native PDF"


# ---------------------------------------------------------------- DOCX
def test_docx_en2ar_translates_text(sample_docx, tmp_path):
    out = str(tmp_path / "out.docx")
    run_pipeline(sample_docx, out, TranslationOptions("en2ar", mirror=True, html_engine=False))
    doc = DocxFile(out)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert contains_arabic(text)
    assert "Quarterly Report" not in text


def test_docx_stores_logical_not_presentation_forms(sample_docx, tmp_path):
    """Word shapes Arabic itself. Storing presentation forms would break search,
    editing and copy-paste."""
    out = str(tmp_path / "out.docx")
    run_pipeline(sample_docx, out, TranslationOptions("en2ar", mirror=True, html_engine=False))
    doc = DocxFile(out)
    for para in doc.paragraphs:
        if para.text.strip():
            assert not is_already_shaped(para.text)


def test_docx_sets_rtl_paragraph_properties(sample_docx, tmp_path):
    out = str(tmp_path / "out.docx")
    run_pipeline(sample_docx, out, TranslationOptions("en2ar", mirror=True, html_engine=False))
    doc = DocxFile(out)
    para = [p for p in doc.paragraphs if p.text.strip()][0]
    pPr = para._p.find(qn("w:pPr"))
    assert pPr is not None and pPr.find(qn("w:bidi")) is not None, \
        "w:bidi is required for correct Arabic paragraph flow"
    assert para.runs[0]._r.find(qn("w:rPr")).find(qn("w:rtl")) is not None


def test_docx_preserves_bold_and_size(sample_docx, tmp_path):
    out = str(tmp_path / "out.docx")
    run_pipeline(sample_docx, out, TranslationOptions("en2ar", mirror=True, html_engine=False))
    doc = DocxFile(out)
    headline = doc.paragraphs[0]
    assert headline.runs[0].bold is True
    assert headline.runs[0].font.size.pt >= 24, "headline size must carry over"


def test_docx_keeps_images(sample_docx, tmp_path):
    import zipfile

    out = str(tmp_path / "out.docx")
    run_pipeline(sample_docx, out, TranslationOptions("en2ar", mirror=True, html_engine=False))
    media = [n for n in zipfile.ZipFile(out).namelist() if "media/" in n]
    assert media, "embedded images must survive the rebuild"


def test_docx_translates_table_cells(sample_docx, tmp_path):
    out = str(tmp_path / "out.docx")
    run_pipeline(sample_docx, out, TranslationOptions("en2ar", mirror=True, html_engine=False))
    doc = DocxFile(out)
    cells = [c.text for row in doc.tables[0].rows for c in row.cells]
    assert any(contains_arabic(c) for c in cells)


def test_docx_underline_can_be_disabled(sample_docx, tmp_path):
    out = str(tmp_path / "out.docx")
    run_pipeline(sample_docx, out,
                 TranslationOptions("en2ar", mirror=True, underline=False, html_engine=False))
    doc = DocxFile(out)
    assert all(not r.underline for p in doc.paragraphs for r in p.runs)


def test_docx_round_trip_returns_english(sample_docx, tmp_path):
    """en2ar then ar2en must give the source text back."""
    arabic = str(tmp_path / "ar.docx")
    english = str(tmp_path / "en.docx")
    run_pipeline(sample_docx, arabic, TranslationOptions("en2ar", mirror=True, html_engine=False))
    run_pipeline(arabic, english, TranslationOptions("ar2en", mirror=True, html_engine=False))

    text = "\n".join(p.text for p in DocxFile(english).paragraphs)
    assert "Quarterly Report" in text
    assert not contains_arabic(text)


def test_docx_output_is_still_docx(sample_docx, tmp_path):
    out = str(tmp_path / "out.docx")
    run_pipeline(sample_docx, out, TranslationOptions("en2ar", html_engine=False))
    with open(out, "rb") as fh:
        assert fh.read(2) == b"PK", "the file must stay a native DOCX"


# ---------------------------------------------------------------- QA
def test_qa_records_font_substitution(sample_pdf, tmp_path):
    out = str(tmp_path / "out.pdf")
    qa = run_pipeline(sample_pdf, out, TranslationOptions("en2ar", mirror=True, html_engine=False))
    report = qa.to_dict()
    assert report["by_category"].get("font_substitution", 0) > 0
    assert any("Noto" in e["message"] for e in report["entries"])


def test_qa_report_is_serialisable(sample_pdf, tmp_path):
    import json

    out = str(tmp_path / "out.pdf")
    qa = run_pipeline(sample_pdf, out, TranslationOptions("en2ar", html_engine=False))
    path = str(tmp_path / "qa.json")
    qa.save(path)
    with open(path, encoding="utf-8") as fh:
        loaded = json.load(fh)
    assert "counts" in loaded and "entries" in loaded


def test_qa_flags_untranslated_segment(sample_pdf, tmp_path):
    """The fake provider leaves one paragraph in English; QA must say so rather
    than degrading silently."""
    out = str(tmp_path / "out.pdf")
    qa = run_pipeline(sample_pdf, out, TranslationOptions("en2ar", mirror=True, html_engine=False))
    messages = [e["message"] for e in qa.to_dict()["entries"]]
    assert any("still in the source language" in m for m in messages)
